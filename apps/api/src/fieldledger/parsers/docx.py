from io import BytesIO

from docx import Document

from fieldledger.domain.models import SourceCoordinate, SourceKind
from fieldledger.parsers.base import ParsedChunk, parsed_chunk


class DocxParser:
    def parse(self, content: bytes) -> tuple[ParsedChunk, ...]:
        document = Document(BytesIO(content))
        chunks: list[ParsedChunk] = []
        for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
            if not paragraph.text.strip():
                continue
            chunks.append(
                parsed_chunk(
                    len(chunks),
                    paragraph.text,
                    SourceCoordinate(
                        kind=SourceKind.DOCX, paragraph=paragraph_number
                    ),
                )
            )
        for table_number, table in enumerate(document.tables, start=1):
            for row_number, row in enumerate(table.rows, start=1):
                for column_number, cell in enumerate(row.cells, start=1):
                    if not cell.text.strip():
                        continue
                    chunks.append(
                        parsed_chunk(
                            len(chunks),
                            cell.text,
                            SourceCoordinate(
                                kind=SourceKind.DOCX,
                                table=table_number,
                                row=row_number,
                                column=column_number,
                            ),
                        )
                    )
        if not chunks:
            raise ValueError("DOCX contains no extractable text")
        return tuple(chunks)
